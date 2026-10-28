"""
Generate professional Word (.docx) and PowerPoint (.pptx) deliverables for Task-1: Comparative Technology Study.
Tailored for ITM SLS BU - School of Computer Science, Engineering & Technology, Vadodara.
Course: Advanced Web Development (B.Tech Semester-5, Winter-2026).
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE

def set_cell_background(cell, fill_hex):
    """Sets background color of a Word table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_word_document(filename):
    doc = Document()
    
    # Page setup - Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # University Header Banner
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_uni = p_uni.add_run("ITM SLS BU\n")
    run_uni.font.name = "Arial"
    run_uni.font.size = Pt(16)
    run_uni.font.bold = True
    run_uni.font.color.rgb = RGBColor(15, 118, 110) # Teal
    
    run_school = p_uni.add_run("School of Computer Science, Engineering & Technology, Vadodara\n")
    run_school.font.name = "Arial"
    run_school.font.size = Pt(12)
    run_school.font.bold = True
    run_school.font.color.rgb = RGBColor(51, 65, 85)

    run_sem = p_uni.add_run("B.Tech Semester-5 | Academic Year 2026-2027, Winter-2026\n")
    run_sem.font.name = "Arial"
    run_sem.font.size = Pt(11)
    run_sem.font.color.rgb = RGBColor(100, 116, 139)

    run_subj = p_uni.add_run("Advanced Web Development (AWD) — Regular Final Viva-Voce Submission\n")
    run_subj.font.name = "Arial"
    run_subj.font.size = Pt(12)
    run_subj.font.bold = True
    run_subj.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Title Box
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t1 = p_title.add_run("TASK-1: COMPARATIVE TECHNOLOGY STUDY\n")
    run_t1.font.name = "Arial"
    run_t1.font.size = Pt(18)
    run_t1.font.bold = True
    run_t1.font.color.rgb = RGBColor(17, 24, 39)

    run_sub = p_title.add_run("CareSync — Full-Stack Healthcare Appointment System (HAS)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Executive Overview
    h_exec = doc.add_heading(level=1)
    r = h_exec.add_run("1. Executive Overview & Scope")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(15, 118, 110)
    
    p = doc.add_paragraph(
        "As part of the Advanced Web Development (AWD) PBL coursework evaluation, this report "
        "delivers an exhaustive comparative analysis of four pivotal technology domains implemented in the "
        "CareSync Healthcare Appointment System. Each study evaluates fundamental concepts, a structured "
        "difference matrix, relative advantages, limitations, suitable industry use cases, and an empirical "
        "justification for why the selected technology was adopted in this project."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)

    # Helper function to add table
    def add_comparison_table(headers, rows):
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Header Row
        hdr_cells = table.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            set_cell_background(hdr_cells[idx], "0F766E")
            set_cell_margins(hdr_cells[idx], top=140, bottom=140, left=160, right=160)
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Arial"
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, row_data in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = cell_value
                set_cell_background(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
                p = row_cells[c_idx].paragraphs[0]
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.bold = True
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(30, 41, 59)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(51, 65, 85)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ------------------ TOPIC 1 ------------------
    h2 = doc.add_heading(level=1)
    r = h2.add_run("2. Comparative Study 1: REST API vs. GraphQL")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(15, 118, 110)

    p = doc.add_paragraph()
    r_sub = p.add_run("2.1 Introduction to Technologies\n")
    r_sub.font.bold = True
    p.add_run(
        "• REST (Representational State Transfer): An established architectural style communicating via standard HTTP methods "
        "(GET, POST, PUT, PATCH, DELETE) across explicit resource endpoints (e.g., /api/appointments, /api/doctors). It relies "
        "on native HTTP status codes (200, 201, 400, 403, 409) and standard cache headers.\n"
        "• GraphQL: A data query language and server runtime developed by Meta that exposes a single endpoint (POST /graphql). "
        "Clients formulate queries declaring exact payload fields, eliminating over-fetching and under-fetching."
    )

    p_t1 = doc.add_paragraph()
    r_t1 = p_t1.add_run("2.2 Key Features & Comparative Difference Matrix")
    r_t1.font.bold = True

    t1_headers = ["Evaluation Criteria", "RESTful API (Express.js)", "GraphQL (Apollo Server)"]
    t1_rows = [
        ["Endpoint Topology", "Multiple distinct resource endpoints (/api/slots, /api/auth)", "Single unified endpoint (POST /graphql)"],
        ["Data Fetching Precision", "Fixed server payload; can suffer from over-fetching", "Client requests exact fields; zero over-fetching"],
        ["Network Caching", "Native HTTP caching via ETags, Cache-Control, and CDN proxies", "Complex; relies on client-side normalized memory cache"],
        ["Error & Status Semantics", "Rich HTTP status codes (200, 201, 400, 403, 409, 429)", "Returns 200 OK for most requests; errors wrapped in JSON array"],
        ["Concurrency & Idempotency", "Standard HTTP verbs, Idempotency-Keys, and 409 Conflict", "Requires custom mutation locks and bespoke conflict resolvers"],
        ["Rate Limiting & Security", "Route-level rate limiting (express-rate-limit), WAF inspection", "Requires AST query depth and complexity cost analysis"],
        ["File Uploads & Streaming", "Standard multipart/form-data supported natively", "Requires schema mutations or multipart spec extensions"]
    ]
    add_comparison_table(t1_headers, t1_rows)

    doc.add_paragraph().add_run("2.3 Advantages & Limitations").font.bold = True
    doc.add_paragraph(
        "• REST Advantages: Standardized HTTP semantics, granular route-level middleware for RBAC, native HTTP proxy/CDN caching, and predictable error handling.\n"
        "• REST Limitations: Fixed response schemas can lead to over-fetching or multiple round-trip requests for deeply nested views.\n"
        "• GraphQL Advantages: Single-request relational fetching, strongly typed schemas, and automatic client code generation.\n"
        "• GraphQL Limitations: Vulnerable to nested recursive DoS queries; lacks native HTTP status codes; complex server caching."
    )

    doc.add_paragraph().add_run("2.4 Suitable Use Cases").font.bold = True
    doc.add_paragraph(
        "• REST API: High-concurrency transactional booking engines, healthcare platforms, microservices, and public enterprise APIs.\n"
        "• GraphQL: Dynamic social feeds, mobile aggregators with variable bandwidth, and complex multi-source dashboards."
    )

    doc.add_paragraph().add_run("2.5 Justified Conclusion for CareSync Project").font.bold = True
    p_just1 = doc.add_paragraph(
        "In CareSync, REST was selected because appointment booking requires strict concurrency control under high contention. "
        "When two patients contend for a single slot, REST responds immediately with standard HTTP 409 Conflict. Furthermore, "
        "granular route middleware enforces RBAC (403 Forbidden) and global rate limiting (429 Too Many Requests) cleanly without "
        "the operational complexity and DoS vulnerabilities of GraphQL query analysis."
    )
    p_just1.paragraph_format.line_spacing = 1.15

    # ------------------ TOPIC 2 ------------------
    h3 = doc.add_heading(level=1)
    r = h3.add_run("3. Comparative Study 2: MongoDB (NoSQL) vs. MySQL (Relational SQL)")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_paragraph().add_run("3.1 Introduction to Technologies").font.bold = True
    doc.add_paragraph(
        "• MongoDB: A document-oriented NoSQL database that stores data as JSON-like BSON documents. Paired with Mongoose ODM in the "
        "MERN stack, it offers rapid schema iteration, atomic single-document operations, and advanced partial indexing.\n"
        "• MySQL: An open-source relational database management system (RDBMS) governed by rigid tables, columns, primary keys, "
        "foreign keys, and ANSI SQL queries, providing strict ACID transactions across normalized relational tables."
    )

    doc.add_paragraph().add_run("3.2 Key Features & Comparative Difference Matrix").font.bold = True
    t2_headers = ["Evaluation Criteria", "MongoDB (Mongoose ODM)", "MySQL (Relational SQL)"]
    t2_rows = [
        ["Data Model", "BSON documents organized in collections", "Rows and columns organized in normalized tables"],
        ["Schema Flexibility", "Dynamic / schema-flexible; validated at ODM level", "Rigid; requires explicit ALTER TABLE DDL migrations"],
        ["Query Language", "Mongoose query builder / MongoDB Aggregation Pipeline", "Structured Query Language (ANSI SQL)"],
        ["Joins vs. Embedding", "Document referencing (.populate()) and embedding", "Multi-table relational JOIN operations"],
        ["Concurrency & Atomicity", "Atomic single-document operations (findOneAndUpdate)", "Row/table locks with multi-table ACID transactions"],
        ["Indexing Features", "Compound indexes, partial filter indexes, TTL indexes", "B-tree primary/foreign key indexes, unique constraints"],
        ["JavaScript Synergy", "Zero impedance mismatch; natively exchanges JSON objects", "Requires ORM (Sequelize/Prisma) mapping layer"]
    ]
    add_comparison_table(t2_headers, t2_rows)

    doc.add_paragraph().add_run("3.3 Advantages & Limitations").font.bold = True
    doc.add_paragraph(
        "• MongoDB Advantages: Natural alignment with JavaScript/Node.js, high throughput atomic compare-and-set operations, "
        "and partial unique indexing for active appointments.\n"
        "• MongoDB Limitations: Complex reporting requires multi-stage aggregation pipelines; higher memory usage for uncompressed schemas.\n"
        "• MySQL Advantages: Enforces mathematical relational integrity and strict multi-table foreign key constraints.\n"
        "• MySQL Limitations: Rigid schema alters introduce migration risks; relational row locking causes bottlenecks under burst slot booking."
    )

    doc.add_paragraph().add_run("3.4 Suitable Use Cases").font.bold = True
    doc.add_paragraph(
        "• MongoDB: MERN stack applications, scheduling platforms, healthcare profiles with variable doctor attributes, and rapid MVP delivery.\n"
        "• MySQL: Core banking, legacy accounting, and highly normalized systems with hundreds of interdependent tables."
    )

    doc.add_paragraph().add_run("3.5 Justified Conclusion for CareSync Project").font.bold = True
    doc.add_paragraph(
        "MongoDB is optimal for CareSync for two vital architectural reasons: First, high-concurrency slot booking requires atomic "
        "compare-and-set execution: 'Slot.findOneAndUpdate({ _id: slotId, isBooked: false }, { $set: { isBooked: true } })'. "
        "This executes in a single atomic database operation without locking other doctors' slots. Second, CareSync leverages MongoDB's "
        "partial unique index on 'Appointment.slotId' with filter '{ status: { $in: [\"PENDING\", \"CONFIRMED\"] } }'. This mathematically "
        "prevents double booking while permitting cancelled/rescheduled appointments to retain audit history without constraint conflicts."
    )

    # ------------------ TOPIC 3 ------------------
    h4 = doc.add_heading(level=1)
    r = h4.add_run("4. Comparative Study 3: React 19 vs. Angular")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_paragraph().add_run("4.1 Introduction to Technologies").font.bold = True
    doc.add_paragraph(
        "• React 19: A component-centric UI library developed by Meta that utilizes declarative JSX syntax, functional components, "
        "Hooks (useState, useEffect, useMemo), unidirectional data flow, and an efficient Virtual DOM reconciler.\n"
        "• Angular: An enterprise-grade, opinionated application framework developed by Google that enforces strict TypeScript, "
        "two-way data binding, RxJS reactive streams, dependency injection, and a monolithic modular architecture."
    )

    doc.add_paragraph().add_run("4.2 Key Features & Comparative Difference Matrix").font.bold = True
    t3_headers = ["Evaluation Criteria", "React 19 (Vite)", "Angular (v18+)"]
    t3_rows = [
        ["Architectural Scope", "UI view library; flexible selection of auxiliary tools", "Batteries-included enterprise MVC/MVVM framework"],
        ["Language & Syntax", "JavaScript / TypeScript with JSX (JavaScript XML)", "Strict TypeScript with HTML templates and decorators"],
        ["Data Flow", "Unidirectional data flow (predictable state down, events up)", "Two-way data binding ([(ngModel)]) and Signals"],
        ["DOM Reconciliation", "Virtual DOM with Fiber reconciler", "Incremental DOM with Zone.js change detection"],
        ["Build & Dev Speed", "Lightning-fast Vite tooling powered by ESBuild / Rollup", "Angular CLI with Webpack/esbuild; heavier compilation"],
        ["Bundle Size", "Extremely lightweight (~45 KB core runtime footprint)", "Heavier baseline framework runtime (~130 KB+)"],
        ["Styling Paradigm", "Seamless integration with Vanilla CSS, CSS Modules, Tailwind", "Component-encapsulated styles (Shadow DOM / Emulated)"]
    ]
    add_comparison_table(t3_headers, t3_rows)

    doc.add_paragraph().add_run("4.3 Advantages & Limitations").font.bold = True
    doc.add_paragraph(
        "• React Advantages: Lightweight, rapid rendering, modular components, minimal boilerplate, and instant hot module replacement.\n"
        "• React Limitations: Unopinionated nature requires manual selection of routing and state libraries; hook dependency mistakes can cause unnecessary renders.\n"
        "• Angular Advantages: Out-of-the-box standardization, built-in dependency injection, and robust form validation.\n"
        "• Angular Limitations: Steep learning curve (RxJS streams, decorators), verbose boilerplate, and heavier runtime footprint."
    )

    doc.add_paragraph().add_run("4.4 Suitable Use Cases").font.bold = True
    doc.add_paragraph(
        "• React: Interactive SPAs, dynamic healthcare booking portals, SaaS dashboards, and consumer-facing web apps.\n"
        "• Angular: Large-scale enterprise intranets and massive monoliths where strict standardization outweighs rapid prototyping."
    )

    doc.add_paragraph().add_run("4.5 Justified Conclusion for CareSync Project").font.bold = True
    doc.add_paragraph(
        "React 19 was selected for CareSync because the platform requires dynamic, responsive UI interactions: a 14-day interactive "
        "slot-picker grid, instantaneous 1-click persona switching (Patient, Doctor, Admin), live unread notification counters, and modal "
        "dialogues. React's declarative state model and Vite build toolchain deliver sub-second hot reloading, fast page loads, and "
        "flawless integration with our custom Vanilla CSS design system."
    )

    # ------------------ TOPIC 4 ------------------
    h5 = doc.add_heading(level=1)
    r = h5.add_run("5. Comparative Study 4: JWT vs. Session-Based Authentication")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_paragraph().add_run("5.1 Introduction to Technologies").font.bold = True
    doc.add_paragraph(
        "• JWT (JSON Web Token): A stateless open standard (RFC 7519) that cryptographically signs claims (userId, role, email) into "
        "a compact base64-encoded token. Transmitted via the 'Authorization: Bearer <token>' header, the server verifies authenticity "
        "cryptographically in CPU memory without database lookups.\n"
        "• Session-Based Authentication: A stateful approach where the server creates a session identifier upon authentication, "
        "stores session state in memory/Redis, and tracks the user via an HTTP-only cookie. Every incoming request requires querying the session store."
    )

    doc.add_paragraph().add_run("5.2 Key Features & Comparative Difference Matrix").font.bold = True
    t4_headers = ["Evaluation Criteria", "JWT (JSON Web Token)", "Session-Based Authentication"]
    t4_rows = [
        ["State Architecture", "100% Stateless; zero server-side memory allocated", "Stateful; server maintains active sessions in Redis/DB"],
        ["Verification Overhead", "CPU cryptographic verification (HMAC-SHA256)", "Cache or database lookup on every single HTTP request"],
        ["Horizontal Scalability", "Effortless; any server instance can verify signed tokens", "Requires centralized cache (Redis) or sticky sessions"],
        ["Client Storage", "Authorization: Bearer header (memory or localStorage)", "HTTP-only, Secure browser cookie"],
        ["CORS & Decoupled Apps", "Trivially supports decoupled origins (Port 5173 to Port 5000)", "Requires complex cookie sharing, SameSite, and domain rules"],
        ["CSRF Attack Vulnerability", "Immune to CSRF when sent in Authorization header", "Vulnerable to CSRF; requires anti-CSRF token handling"],
        ["Revocation Control", "Challenging before expiry (requires blacklist/TTL)", "Instant; server simply purges session key from Redis"]
    ]
    add_comparison_table(t4_headers, t4_rows)

    doc.add_paragraph().add_run("5.3 Advantages & Limitations").font.bold = True
    doc.add_paragraph(
        "• JWT Advantages: Stateless scalability, zero database overhead for auth verification, decoupled SPA compatibility, and built-in role claims for RBAC.\n"
        "• JWT Limitations: Token cannot be easily invalidated before expiration; payload claims are visible if not encrypted.\n"
        "• Session Advantages: Instant administrative revocation and smaller cookie payload.\n"
        "• Session Limitations: High database/cache lookup overhead; scaling bottlenecks; CSRF vulnerabilities."
    )

    doc.add_paragraph().add_run("5.4 Suitable Use Cases").font.bold = True
    doc.add_paragraph(
        "• JWT: Decoupled RESTful APIs, modern SPAs (React + Express), mobile apps, and microservices.\n"
        "• Session: Monolithic server-side rendered applications (Express with EJS, PHP, traditional Django) hosted under a single root domain."
    )

    doc.add_paragraph().add_run("5.5 Justified Conclusion for CareSync Project").font.bold = True
    doc.add_paragraph(
        "JWT-based authentication is the ideal choice for CareSync because the React client and Express API run as decoupled services. "
        "By embedding user role claims (PATIENT, DOCTOR, ADMIN) directly into the token payload, our RBAC middleware enforces route permissions "
        "instantly in memory without issuing redundant MongoDB queries on every request. Short-lived tokens combined with client-side cleanup "
        "on logout provide an optimal balance of high performance, security, and developer velocity."
    )

    # ------------------ SUMMARY ------------------
    h6 = doc.add_heading(level=1)
    r = h6.add_run("6. Architectural Synthesis & Conclusion")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_paragraph(
        "The four selected technologies — RESTful API, MongoDB, React 19, and JWT Authentication — form a synergistic, high-performance "
        "MERN architecture specifically designed to solve healthcare scheduling challenges:\n\n"
        "1. Concurrency Safety: Atomic MongoDB operations prevent double booking.\n"
        "2. Performance & Low Latency: Stateless JWTs and lightweight React components deliver sub-second user interactions.\n"
        "3. Security & Access Control: Granular REST middleware enforces role-based access boundaries.\n"
        "4. Academic & Industry Standard: The architecture fully adheres to the B.Tech Semester-5 AWD syllabus and PBL guidelines."
    )

    doc.save(filename)
    print(f"[SUCCESS] Word document generated: {filename}")

def create_powerpoint_presentation(filename):
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    blank_layout = prs.slide_layouts[6]

    DARK_TEAL = PptRGBColor(15, 118, 110)
    DARK_BLUE = PptRGBColor(15, 23, 42)
    LIGHT_BG = PptRGBColor(248, 250, 252)
    WHITE = PptRGBColor(255, 255, 255)
    GRAY_TEXT = PptRGBColor(71, 85, 105)
    ACCENT_CYAN = PptRGBColor(13, 148, 136)

    def add_header(slide, title_text, category_text="TASK-1: COMPARATIVE TECHNOLOGY STUDY"):
        # Header background banner
        header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(13.333), PptInches(1.2))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = DARK_BLUE
        header_shape.line.fill.background()

        # Category
        tx_cat = slide.shapes.add_textbox(PptInches(0.8), PptInches(0.15), PptInches(11.5), PptInches(0.3))
        tf_c = tx_cat.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.text = f"ITM SLS BU — AWD SEMESTER-5 | {category_text}"
        p_c.font.size = PptPt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = ACCENT_CYAN

        # Title
        tx_title = slide.shapes.add_textbox(PptInches(0.8), PptInches(0.45), PptInches(11.5), PptInches(0.6))
        tf_t = tx_title.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = PptPt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

    # ------------------ SLIDE 1: Title Slide ------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(13.333), PptInches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BLUE
    bg1.line.fill.background()

    tx = slide1.shapes.add_textbox(PptInches(1.0), PptInches(1.2), PptInches(11.333), PptInches(5.0))
    tf = tx.text_frame
    
    p = tf.paragraphs[0]
    p.text = "ITM SLS BU — School of Computer Science, Engineering & Technology"
    p.font.size = PptPt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = PptPt(14)

    p2 = tf.add_paragraph()
    p2.text = "TASK-1: COMPARATIVE TECHNOLOGY STUDY"
    p2.font.size = PptPt(34)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = PptPt(10)

    p3 = tf.add_paragraph()
    p3.text = "CareSync — Healthcare Appointment System (HAS)"
    p3.font.size = PptPt(22)
    p3.font.bold = True
    p3.font.color.rgb = PptRGBColor(204, 251, 241)
    p3.space_after = PptPt(24)

    p4 = tf.add_paragraph()
    p4.text = (
        "Course: Advanced Web Development (AWD) | B.Tech Semester-5 (Winter-2026)\n"
        "Evaluation: Regular Final Viva-Voce Examination\n"
        "Topics Covered: REST vs. GraphQL | MongoDB vs. MySQL | React vs. Angular | JWT vs. Sessions"
    )
    p4.font.size = PptPt(14)
    p4.font.color.rgb = PptRGBColor(148, 163, 184)

    # ------------------ SLIDE 2: Agenda & Architecture ------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Executive Overview: Core Architectural Comparisons")

    # 4 Cards
    cards = [
        ("1. API Layer", "RESTful API vs. GraphQL", "Resource endpoints, HTTP semantics, granular rate limiting, and conflict resolution (409 Conflict).", "Selected: RESTful Express API"),
        ("2. Database Engine", "MongoDB vs. MySQL", "Document flexibility, atomic compare-and-set slot claiming, and partial unique indexes.", "Selected: MongoDB + Mongoose"),
        ("3. Client Architecture", "React 19 vs. Angular", "Component composition, reactive slot picker, sub-second Vite HMR, and lightweight bundle.", "Selected: React 19 + Vite"),
        ("4. Authentication", "JWT vs. Session Auth", "Stateless bearer authorization, instant RBAC in memory, and decoupled cross-origin frontend.", "Selected: JWT Bearer Tokens")
    ]

    for idx, (cat, title, desc, sel) in enumerate(cards):
        left = PptInches(1.0 + (idx % 2) * 5.8)
        top = PptInches(1.8 + (idx // 2) * 2.6)
        
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, PptInches(5.5), PptInches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = PptRGBColor(226, 232, 240)
        
        tf_c = card.text_frame
        p_c1 = tf_c.paragraphs[0]
        p_c1.text = cat.upper()
        p_c1.font.size = PptPt(11)
        p_c1.font.bold = True
        p_c1.font.color.rgb = DARK_TEAL
        
        p_c2 = tf_c.add_paragraph()
        p_c2.text = title
        p_c2.font.size = PptPt(16)
        p_c2.font.bold = True
        p_c2.font.color.rgb = DARK_BLUE
        
        p_c3 = tf_c.add_paragraph()
        p_c3.text = desc
        p_c3.font.size = PptPt(11)
        p_c3.font.color.rgb = GRAY_TEXT
        
        p_c4 = tf_c.add_paragraph()
        p_c4.text = f"✔ {sel}"
        p_c4.font.size = PptPt(11.5)
        p_c4.font.bold = True
        p_c4.font.color.rgb = DARK_TEAL

    # ------------------ SLIDE 3: REST vs GraphQL ------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Comparative Study 1: RESTful API vs. GraphQL")

    table_shape = slide3.shapes.add_table(5, 3, PptInches(0.8), PptInches(1.6), PptInches(11.733), PptInches(3.2))
    tbl = table_shape.table
    tbl.columns[0].width = PptInches(2.5)
    tbl.columns[1].width = PptInches(4.6)
    tbl.columns[2].width = PptInches(4.6)

    t_headers = ["Dimension", "RESTful API (CareSync)", "GraphQL (Alternative)"]
    for i, h in enumerate(t_headers):
        cell = tbl.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_TEAL
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.size = PptPt(12)

    s3_data = [
        ("Endpoint Structure", "Multiple resource endpoints (/api/appointments, /api/slots)", "Single endpoint (POST /graphql) with query payload"),
        ("HTTP Status Codes", "Rich semantics: 200, 201, 400, 403, 409, 429", "Returns 200 OK almost always; errors in payload array"),
        ("Concurrency & Locks", "Standard 409 Conflict on double booking; idempotency headers", "Requires custom mutation resolvers and lock management"),
        ("Caching & WAF", "Native HTTP caching, CDN integration, route-level rate limiter", "Complex query cost analysis; difficult CDN edge caching")
    ]
    for row_idx, row_vals in enumerate(s3_data):
        for col_idx, val in enumerate(row_vals):
            cell = tbl.cell(row_idx + 1, col_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = PptPt(11)
            p.font.color.rgb = DARK_BLUE if col_idx == 0 else GRAY_TEXT
            if col_idx == 0:
                p.font.bold = True

    # Justification Banner
    just_card = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.8), PptInches(5.2), PptInches(11.733), PptInches(1.8))
    just_card.fill.solid()
    just_card.fill.fore_color.rgb = PptRGBColor(240, 253, 250)
    just_card.line.color.rgb = DARK_TEAL
    tf_j = just_card.text_frame
    p_jh = tf_j.paragraphs[0]
    p_jh.text = "PROJECT JUSTIFICATION FOR CARESYNC:"
    p_jh.font.bold = True
    p_jh.font.size = PptPt(12)
    p_jh.font.color.rgb = DARK_TEAL
    
    p_jd = tf_j.add_paragraph()
    p_jd.text = (
        "• High Concurrency: CareSync handles slot contention by returning 409 Conflict when a slot is already taken.\n"
        "• Strict RBAC Boundaries: Granular Express middleware guards Patient, Doctor, and Admin routes independently.\n"
        "• Predictable Security: Global & endpoint rate limiters prevent DoS without needing complex AST query parsing."
    )
    p_jd.font.size = PptPt(11)
    p_jd.font.color.rgb = DARK_BLUE

    # ------------------ SLIDE 4: MongoDB vs MySQL ------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Comparative Study 2: MongoDB (NoSQL) vs. MySQL (SQL)")

    table_shape = slide4.shapes.add_table(5, 3, PptInches(0.8), PptInches(1.6), PptInches(11.733), PptInches(3.2))
    tbl = table_shape.table
    tbl.columns[0].width = PptInches(2.5)
    tbl.columns[1].width = PptInches(4.6)
    tbl.columns[2].width = PptInches(4.6)

    t_headers = ["Dimension", "MongoDB (Mongoose ODM)", "MySQL (Relational SQL)"]
    for i, h in enumerate(t_headers):
        cell = tbl.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_TEAL
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.size = PptPt(12)

    s4_data = [
        ("Data Model", "BSON documents; dynamic schema with Mongoose validation", "Rigid tabular rows and columns; fixed DDL schemas"),
        ("Atomic Operations", "Atomic findOneAndUpdate compare-and-set slot claiming", "Row/table locking (InnoDB); requires transaction blocks"),
        ("Advanced Indexing", "Compound unique (doctorId, date, time) + Partial Unique Index", "B-tree primary/foreign key indexes; triggers for conditional unique"),
        ("Node.js Synergy", "Native JSON/JavaScript integration; zero ORM impedance mismatch", "Requires ORM translation layer (Sequelize/Prisma)")
    ]
    for row_idx, row_vals in enumerate(s4_data):
        for col_idx, val in enumerate(row_vals):
            cell = tbl.cell(row_idx + 1, col_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = PptPt(11)
            p.font.color.rgb = DARK_BLUE if col_idx == 0 else GRAY_TEXT
            if col_idx == 0:
                p.font.bold = True

    # Justification Banner
    just_card4 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.8), PptInches(5.2), PptInches(11.733), PptInches(1.8))
    just_card4.fill.solid()
    just_card4.fill.fore_color.rgb = PptRGBColor(240, 253, 250)
    just_card4.line.color.rgb = DARK_TEAL
    tf_j = just_card4.text_frame
    p_jh = tf_j.paragraphs[0]
    p_jh.text = "PROJECT JUSTIFICATION FOR CARESYNC:"
    p_jh.font.bold = True
    p_jh.font.size = PptPt(12)
    p_jh.font.color.rgb = DARK_TEAL
    
    p_jd = tf_j.add_paragraph()
    p_jd.text = (
        "• Atomic Compare-and-Set: 'Slot.findOneAndUpdate({ _id: slotId, isBooked: false }, { $set: { isBooked: true } })' eliminates double-booking in a single DB operation.\n"
        "• Partial Unique Indexing: 'appointmentSchema.index({ slotId: 1 }, { unique: true, partialFilterExpression: { status: { $in: [\"PENDING\", \"CONFIRMED\"] } } })'.\n"
        "• Flexible Clinical Profiles: Seamlessly handles diverse doctor qualifications, bio, and weekly availability schedules."
    )
    p_jd.font.size = PptPt(11)
    p_jd.font.color.rgb = DARK_BLUE

    # ------------------ SLIDE 5: React vs Angular ------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Comparative Study 3: React 19 vs. Angular")

    table_shape = slide5.shapes.add_table(5, 3, PptInches(0.8), PptInches(1.6), PptInches(11.733), PptInches(3.2))
    tbl = table_shape.table
    tbl.columns[0].width = PptInches(2.5)
    tbl.columns[1].width = PptInches(4.6)
    tbl.columns[2].width = PptInches(4.6)

    t_headers = ["Dimension", "React 19 (Vite)", "Angular (v18+)"]
    for i, h in enumerate(t_headers):
        cell = tbl.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_TEAL
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.size = PptPt(12)

    s5_data = [
        ("Architecture", "Declarative UI component library; modular ecosystem", "Monolithic enterprise MVC/MVVM framework"),
        ("Language & Syntax", "JavaScript/TypeScript with JSX; functional components + Hooks", "Strict TypeScript with HTML templates, decorators & RxJS"),
        ("Build & Dev Speed", "Vite build toolchain (ESBuild); millisecond HMR", "Angular CLI with Webpack/esbuild; heavier compilation overhead"),
        ("Bundle Size", "Lightweight (~45 KB core runtime footprint)", "Heavier baseline framework runtime (~130 KB+ baseline)")
    ]
    for row_idx, row_vals in enumerate(s5_data):
        for col_idx, val in enumerate(row_vals):
            cell = tbl.cell(row_idx + 1, col_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = PptPt(11)
            p.font.color.rgb = DARK_BLUE if col_idx == 0 else GRAY_TEXT
            if col_idx == 0:
                p.font.bold = True

    # Justification Banner
    just_card5 = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.8), PptInches(5.2), PptInches(11.733), PptInches(1.8))
    just_card5.fill.solid()
    just_card5.fill.fore_color.rgb = PptRGBColor(240, 253, 250)
    just_card5.line.color.rgb = DARK_TEAL
    tf_j = just_card5.text_frame
    p_jh = tf_j.paragraphs[0]
    p_jh.text = "PROJECT JUSTIFICATION FOR CARESYNC:"
    p_jh.font.bold = True
    p_jh.font.size = PptPt(12)
    p_jh.font.color.rgb = DARK_TEAL
    
    p_jd = tf_j.add_paragraph()
    p_jd.text = (
        "• Fluid Multi-Role Experience: Seamless 1-click persona switching (Patient, Doctor, Admin) without full-page reloads.\n"
        "• Interactive 14-Day Slot Grid: Declarative state updates render available time slots in real time.\n"
        "• Pure Vanilla CSS System: Allows direct, unencumbered CSS styling for modern glassmorphism, responsive cards, and zero UI lag."
    )
    p_jd.font.size = PptPt(11)
    p_jd.font.color.rgb = DARK_BLUE

    # ------------------ SLIDE 6: JWT vs Sessions ------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Comparative Study 4: JWT vs. Session-Based Authentication")

    table_shape = slide6.shapes.add_table(5, 3, PptInches(0.8), PptInches(1.6), PptInches(11.733), PptInches(3.2))
    tbl = table_shape.table
    tbl.columns[0].width = PptInches(2.5)
    tbl.columns[1].width = PptInches(4.6)
    tbl.columns[2].width = PptInches(4.6)

    t_headers = ["Dimension", "JWT Authentication (CareSync)", "Session Authentication (Alternative)"]
    for i, h in enumerate(t_headers):
        cell = tbl.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_TEAL
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.size = PptPt(12)

    s6_data = [
        ("State Management", "100% Stateless; zero server memory allocated", "Stateful; server must maintain active session records in Redis/DB"),
        ("Verification Speed", "Cryptographic HMAC-SHA256 verification in CPU memory", "Database/Cache round-trip required on every single request"),
        ("Horizontal Scale", "Trivially scalable; any server instance validates tokens", "Requires centralized cache (Redis) or sticky sessions"),
        ("CORS & SPAs", "Authorization: Bearer header natively handles decoupled ports", "Complex cookie sharing, SameSite flags, and CSRF vulnerabilities")
    ]
    for row_idx, row_vals in enumerate(s6_data):
        for col_idx, val in enumerate(row_vals):
            cell = tbl.cell(row_idx + 1, col_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = PptPt(11)
            p.font.color.rgb = DARK_BLUE if col_idx == 0 else GRAY_TEXT
            if col_idx == 0:
                p.font.bold = True

    # Justification Banner
    just_card6 = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.8), PptInches(5.2), PptInches(11.733), PptInches(1.8))
    just_card6.fill.solid()
    just_card6.fill.fore_color.rgb = PptRGBColor(240, 253, 250)
    just_card6.line.color.rgb = DARK_TEAL
    tf_j = just_card6.text_frame
    p_jh = tf_j.paragraphs[0]
    p_jh.text = "PROJECT JUSTIFICATION FOR CARESYNC:"
    p_jh.font.bold = True
    p_jh.font.size = PptPt(12)
    p_jh.font.color.rgb = DARK_TEAL
    
    p_jd = tf_j.add_paragraph()
    p_jd.text = (
        "• In-Memory RBAC Validation: Role claims (PATIENT, DOCTOR, ADMIN) are verified instantly without querying MongoDB.\n"
        "• Decoupled Architecture: React on Port 5173 communicates with Express on Port 5000 without third-party cookie restrictions.\n"
        "• 1-Click Demo Evaluation: Tokens are issued immediately for seamless viva evaluation and multi-role testing."
    )
    p_jd.font.size = PptPt(11)
    p_jd.font.color.rgb = DARK_BLUE

    # ------------------ SLIDE 7: Conclusion & Synergy ------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Architectural Synthesis & Conclusion")

    tx7 = slide7.shapes.add_textbox(PptInches(1.0), PptInches(1.6), PptInches(11.333), PptInches(4.5))
    tf7 = tx7.text_frame

    p = tf7.paragraphs[0]
    p.text = "Unified Full-Stack Synergy in CareSync:"
    p.font.size = PptPt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_TEAL
    p.space_after = PptPt(16)

    bullets = [
        ("High Concurrency Protection: ", "MongoDB atomic compare-and-set + partial unique indexes prevent double-booking at the database engine level."),
        ("Predictable API Semantics: ", "Express REST API provides RFC-compliant HTTP status codes (201, 403, 409, 429) and standard response envelopes."),
        ("Instantaneous Access Control: ", "Stateless JWT Bearer tokens provide sub-millisecond RBAC route authorization in memory."),
        ("Modern User Experience: ", "React 19 + Vite delivers responsive slot picking, instant role switching, and zero UI latency with pure Vanilla CSS."),
        ("Full Curriculum Compliance: ", "Strictly conforms to ITM SLS BU Semester-5 AWD requirements and Project-Based Learning objectives.")
    ]

    for title, desc in bullets:
        p_b = tf7.add_paragraph()
        p_b.space_after = PptPt(10)
        r_t = p_b.add_run()
        r_t.text = f"• {title}"
        r_t.font.bold = True
        r_t.font.size = PptPt(13)
        r_t.font.color.rgb = DARK_BLUE
        
        r_d = p_b.add_run()
        r_d.text = desc
        r_d.font.size = PptPt(13)
        r_d.font.color.rgb = GRAY_TEXT

    prs.save(filename)
    print(f"[SUCCESS] PowerPoint presentation generated: {filename}")

if __name__ == "__main__":
    docs_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(docs_dir, "4631_Task_1_Comparative_Technology_Study.docx")
    pptx_path = os.path.join(docs_dir, "4631_Task_1_Comparative_Technology_Study.pptx")
    
    create_word_document(docx_path)
    create_powerpoint_presentation(pptx_path)
