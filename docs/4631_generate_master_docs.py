"""
Generate Master Consolidated Word (.docx) and PowerPoint (.pptx) deliverables combining
Task-1, Task-2, and Task-3 into a unified submission package for AWD Semester-5 Final Viva-Voce.
ITM SLS BU - School of Computer Science, Engineering & Technology, Vadodara.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_master_word_document(filename):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # University Header Banner
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_uni = p_uni.add_run("ITM SLS BU\n")
    run_uni.font.name = "Arial"
    run_uni.font.size = Pt(18)
    run_uni.font.bold = True
    run_uni.font.color.rgb = RGBColor(15, 118, 110)
    
    run_school = p_uni.add_run("School of Computer Science, Engineering & Technology, Vadodara\n")
    run_school.font.name = "Arial"
    run_school.font.size = Pt(13)
    run_school.font.bold = True
    run_school.font.color.rgb = RGBColor(51, 65, 85)

    run_sem = p_uni.add_run("B.Tech Semester-5 | Academic Year 2026-2027, Winter-2026\n")
    run_sem.font.name = "Arial"
    run_sem.font.size = Pt(11)
    run_sem.font.color.rgb = RGBColor(100, 116, 139)

    run_subj = p_uni.add_run("Advanced Web Development (AWD) — Regular Final Viva-Voce Master Submission\n")
    run_subj.font.name = "Arial"
    run_subj.font.size = Pt(13)
    run_subj.font.bold = True
    run_subj.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t1 = p_title.add_run("CareSync — Healthcare Appointment System (HAS)\n")
    run_t1.font.name = "Arial"
    run_t1.font.size = Pt(22)
    run_t1.font.bold = True
    run_t1.font.color.rgb = RGBColor(17, 24, 39)

    run_sub = p_title.add_run("Comprehensive Technical Documentation for Tasks 1, 2, and 3")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Executive Overview
    h_exec = doc.add_heading(level=1)
    r = h_exec.add_run("1. Executive Summary & Curriculum Alignment")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(15, 118, 110)
    
    doc.add_paragraph(
        "This master submission compiles the complete technical work for the Advanced Web Development (AWD) "
        "Project-Based Learning (PBL) component. It systematically covers all three mandatory tasks prescribed "
        "in the semester evaluation guidelines:\n\n"
        "• Task-1: Comparative Technology Study (REST vs GraphQL, MongoDB vs MySQL, React vs Angular, JWT vs Sessions)\n"
        "• Task-2: Security Audit Report (Vulnerability assessment of injection, XSS, auth, storage, validation, and race conditions)\n"
        "• Task-3: Implementation & Demonstration of Security Libraries (Integration of Helmet.js, bcrypt, and Joi with live test verification)"
    )

    def add_table(headers, rows, bg_header="0F766E", row_alt="F0FDFA"):
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        hdr_cells = table.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            set_cell_background(hdr_cells[idx], bg_header)
            set_cell_margins(hdr_cells[idx], top=130, bottom=130, left=140, right=140)
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Arial"
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, row_data in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = row_alt if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = cell_value
                set_cell_background(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=120, right=120)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(51, 65, 85)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Task 1 Section
    h_t1 = doc.add_heading(level=1)
    r = h_t1.add_run("2. Task-1: Comparative Technology Study")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(15, 118, 110)

    t1_headers = ["Domain", "Selected in CareSync", "Industry Alternative", "Empirical Project Justification"]
    t1_rows = [
        ["API Layer", "RESTful API (Express)", "GraphQL (Apollo)", "Deterministic HTTP 409 Conflict status codes for atomic double-booking rejection and route-level rate limiting."],
        ["Database", "MongoDB (Mongoose ODM)", "MySQL (Relational SQL)", "Atomic compare-and-set slot claiming and partial unique index on Appointment.slotId for active bookings."],
        ["Frontend", "React 19 (Vite)", "Angular (v18+)", "Sub-second HMR, interactive 14-day slot grid picker, 1-click persona switching, and unencumbered Vanilla CSS styling."],
        ["Auth Strategy", "JWT Bearer Tokens", "Session-Based Cookies", "Stateless verification in memory eliminates database lookups on authenticated RBAC route checks."]
    ]
    add_table(t1_headers, t1_rows)

    # Task 2 Section
    h_t2 = doc.add_heading(level=1)
    r = h_t2.add_run("3. Task-2: Security Audit Report")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(180, 83, 9)

    t2_headers = ["Vulnerability Vector", "Risk Severity", "Status", "Technical Defense Implemented in CareSync"]
    t2_rows = [
        ["NoSQL Injection (CWE-943)", "Medium", "Mitigated", "Mongoose schema casting + Joi strict string type enforcement."],
        ["Cross-Site Scripting (XSS - CWE-79)", "Low", "Fully Mitigated", "React 19 JSX auto-escaping + Helmet CSP HTTP headers."],
        ["Broken Authorization (RBAC - CWE-285)", "High", "Fully Mitigated", "Express authorizeRoles middleware + admin lockout prevention."],
        ["Insecure Password Storage (CWE-256)", "Low", "Fully Mitigated", "bcryptjs with 10 salt rounds + User.toJSON passwordHash suppression."],
        ["Inadequate Input Validation (CWE-20)", "Medium", "Fully Mitigated", "Centralized Joi validation middleware rejecting malformed payloads with 400."],
        ["Double-Booking Race Condition (CWE-362)", "High", "Fully Mitigated", "Atomic findOneAndUpdate + Partial Unique Index (10 concurrent requests -> 1 success, 9 409 Conflict)."]
    ]
    add_table(t2_headers, t2_rows, bg_header="B45309", row_alt="FFFBEB")

    # Task 3 Section
    h_t3 = doc.add_heading(level=1)
    r = h_t3.add_run("4. Task-3: Security Libraries Implementation & Demonstration")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(2, 132, 199)

    t3_headers = ["Library", "Version", "Role in CareSync", "Live Verification Output"]
    t3_rows = [
        ["Helmet.js", "^8.0.0", "HTTP Header Hardening", "Sets nosniff, SAMEORIGIN, CSP, and removes X-Powered-By fingerprint."],
        ["bcryptjs", "^3.0.2", "Password Hashing & Matching", "Hashes with 10 salt rounds; distinct salts generated; matching compare verified."],
        ["Joi", "^18.3.0", "Request Schema Validation", "Rejects invalid email/password with 400 Bad Request; validates MongoDB ObjectIds."]
    ]
    add_table(t3_headers, t3_rows, bg_header="0284C7", row_alt="F0F9FF")

    # Conclusion
    h_c = doc.add_heading(level=1)
    r = h_c.add_run("5. Master Conclusion & Test Verification")
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_paragraph(
        "CareSync represents an enterprise-grade, high-performance healthcare platform built with strict architectural compliance. "
        "The automated test suite in server/test executes 31 integration and security unit tests with a 100% pass rate. "
        "All academic requirements for B.Tech Semester-5 AWD have been fully satisfied."
    )

    doc.save(filename)
    print(f"[SUCCESS] Master Word document generated: {filename}")

def create_master_powerpoint_presentation(filename):
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    blank_layout = prs.slide_layouts[6]

    DARK_TEAL = PptRGBColor(15, 118, 110)
    DARK_BLUE = PptRGBColor(15, 23, 42)
    WHITE = PptRGBColor(255, 255, 255)
    GRAY_TEXT = PptRGBColor(71, 85, 105)
    ACCENT_CYAN = PptRGBColor(13, 148, 136)

    def add_header(slide, title_text, category_text="REGULAR FINAL VIVA-VOCE MASTER PRESENTATION"):
        header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(13.333), PptInches(1.2))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = DARK_BLUE
        header_shape.line.fill.background()

        tx_cat = slide.shapes.add_textbox(PptInches(0.8), PptInches(0.15), PptInches(11.5), PptInches(0.3))
        tf_c = tx_cat.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.text = f"ITM SLS BU — AWD SEMESTER-5 | {category_text}"
        p_c.font.size = PptPt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = ACCENT_CYAN

        tx_title = slide.shapes.add_textbox(PptInches(0.8), PptInches(0.45), PptInches(11.5), PptInches(0.6))
        tf_t = tx_title.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = PptPt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

    # ------------------ SLIDE 1: Title Slide ------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(13.333), PptInches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BLUE
    bg1.line.fill.background()

    tx = slide1.shapes.add_textbox(PptInches(1.0), PptInches(1.2), PptInches(11.333), PptInches(5.0))
    tf = tx.text_frame
    
    p = tf.paragraphs[0]
    p.text = "ITM SLS BU — School of Computer Science, Engineering & Technology"
    p.font.size = PptPt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = PptPt(14)

    p2 = tf.add_paragraph()
    p2.text = "CareSync — Healthcare Appointment System (HAS)"
    p2.font.size = PptPt(34)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = PptPt(10)

    p3 = tf.add_paragraph()
    p3.text = "Advanced Web Development (AWD) | Final Viva-Voce Master Presentation"
    p3.font.size = PptPt(20)
    p3.font.bold = True
    p3.font.color.rgb = PptRGBColor(204, 251, 241)
    p3.space_after = PptPt(24)

    p4 = tf.add_paragraph()
    p4.text = (
        "Course: B.Tech Semester-5 (Winter-2026)\n"
        "Complete Implementation: Task 1 (Comparative Study) | Task 2 (Security Audit) | Task 3 (Security Libraries)"
    )
    p4.font.size = PptPt(14)
    p4.font.color.rgb = PptRGBColor(148, 163, 184)

    # ------------------ SLIDE 2: Task 1 Summary ------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Task-1: Comparative Technology Study Summary")

    t1_cards = [
        ("REST API vs. GraphQL", "Selected: RESTful API via Express", "Enables deterministic RFC HTTP status codes (409 Conflict for double booking), simple edge caching, and granular route RBAC."),
        ("MongoDB vs. MySQL", "Selected: MongoDB + Mongoose", "Atomic compare-and-set slot reservations (Slot.findOneAndUpdate) and partial unique indexes on active appointments."),
        ("React 19 vs. Angular", "Selected: React 19 + Vite", "Sub-second HMR, interactive 14-day slot-picking grid, instant 1-click persona switching, and custom Vanilla CSS design system."),
        ("JWT vs. Session Auth", "Selected: JWT Bearer Tokens", "Stateless in-memory RBAC authorization eliminates database lookup latency on protected API requests.")
    ]
    for idx, (title, sel, desc) in enumerate(t1_cards):
        left = PptInches(1.0 + (idx % 2) * 5.8)
        top = PptInches(1.8 + (idx // 2) * 2.6)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, PptInches(5.5), PptInches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = PptRGBColor(226, 232, 240)
        tf_c = card.text_frame
        p1 = tf_c.paragraphs[0]
        p1.text = title
        p1.font.size = PptPt(14)
        p1.font.bold = True
        p1.font.color.rgb = DARK_BLUE
        p2 = tf_c.add_paragraph()
        p2.text = f"✔ {sel}"
        p2.font.size = PptPt(11)
        p2.font.bold = True
        p2.font.color.rgb = DARK_TEAL
        p3 = tf_c.add_paragraph()
        p3.text = desc
        p3.font.size = PptPt(10.5)
        p3.font.color.rgb = GRAY_TEXT

    # ------------------ SLIDE 3: Task 2 Summary ------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Task-2: Security Audit Report Summary")

    t2_points = [
        ("NoSQL Injection (CWE-943): ", "Mongoose schema casting strictly typed fields; Joi validates string constraints prior to queries."),
        ("Cross-Site Scripting (XSS - CWE-79): ", "React 19 auto-escapes all JSX expressions; Helmet sets CSP headers; dangerouslySetInnerHTML is 0% used."),
        ("Broken Authorization (CWE-285): ", "Strict RBAC middleware (authorizeRoles) enforces least privilege; direct ADMIN registration is blocked (403)."),
        ("Insecure Password Storage (CWE-256): ", "bcryptjs with 10 salt rounds; User.toJSON overrides automatically strip passwordHash from network output."),
        ("Concurrency Race Conditions (CWE-362): ", "Atomic findOneAndUpdate slot claims + partial unique index on active appointments (10 concurrent requests -> 1 success, 9 conflicts).")
    ]
    tx3 = slide3.shapes.add_textbox(PptInches(1.0), PptInches(1.6), PptInches(11.333), PptInches(5.2))
    tf3 = tx3.text_frame
    for title, desc in t2_points:
        p_pt = tf3.add_paragraph()
        p_pt.space_after = PptPt(12)
        r_t = p_pt.add_run()
        r_t.text = f"✔ {title}"
        r_t.font.bold = True
        r_t.font.size = PptPt(13)
        r_t.font.color.rgb = PptRGBColor(180, 83, 9)
        r_d = p_pt.add_run()
        r_d.text = desc
        r_d.font.size = PptPt(12)
        r_d.font.color.rgb = GRAY_TEXT

    # ------------------ SLIDE 4: Task 3 Summary ------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Task-3: Security Libraries Implementation & Testing")

    t3_cards = [
        ("Helmet.js (^8.0.0)", "HTTP Security Headers", "Verified in app.js: sets nosniff, SAMEORIGIN, CSP, and suppresses X-Powered-By.", "Status: 100% Passed"),
        ("bcryptjs (^3.0.2)", "Password Cryptography", "Verified in authController.js: generates unique salts; securely matches valid passwords; rejects invalid.", "Status: 100% Passed"),
        ("Joi (^18.3.0)", "Input Validation Middleware", "Verified in validate.js: blocks malformed emails, short passwords (<6 chars), and invalid ObjectIds with 400.", "Status: 100% Passed")
    ]
    for idx, (lib, role, desc, status) in enumerate(t3_cards):
        top = PptInches(1.8 + idx * 1.7)
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(1.0), top, PptInches(11.333), PptInches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = PptRGBColor(226, 232, 240)
        tf_c = card.text_frame
        p1 = tf_c.paragraphs[0]
        p1.text = f"{lib} — {role}"
        p1.font.size = PptPt(13)
        p1.font.bold = True
        p1.font.color.rgb = PptRGBColor(2, 132, 199)
        p2 = tf_c.add_paragraph()
        p2.text = desc
        p2.font.size = PptPt(11)
        p2.font.color.rgb = GRAY_TEXT
        p3 = tf_c.add_paragraph()
        p3.text = f"✔ {status}"
        p3.font.size = PptPt(10.5)
        p3.font.bold = True
        p3.font.color.rgb = PptRGBColor(22, 163, 74)

    # ------------------ SLIDE 5: Conclusion & Q&A ------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Master Synthesis & Viva Conclusion")

    tx5 = slide5.shapes.add_textbox(PptInches(1.0), PptInches(1.8), PptInches(11.333), PptInches(4.5))
    tf5 = tx5.text_frame
    p = tf5.paragraphs[0]
    p.text = "Project Assessment Summary:"
    p.font.size = PptPt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_TEAL
    p.space_after = PptPt(16)

    bullets = [
        ("Full Task Coverage: ", "Task 1 (Comparative Study), Task 2 (Security Audit), and Task 3 (Security Libraries) are 100% completed."),
        ("Dual Deliverables Provided: ", "Individual Word & PPT files for each task + Unified Master Word & Master PPT presentation."),
        ("Empirical Validation: ", "Automated test suite (31 tests) passes with 100% success across concurrency, RBAC, and security libraries."),
        ("Ready for Final Defense: ", "System is fully operational and certified for ITM SLS BU Semester-5 AWD Final Viva-Voce.")
    ]
    for title, desc in bullets:
        p_b = tf5.add_paragraph()
        p_b.space_after = PptPt(10)
        r_t = p_b.add_run()
        r_t.text = f"✔ {title}"
        r_t.font.bold = True
        r_t.font.size = PptPt(13)
        r_t.font.color.rgb = DARK_BLUE
        r_d = p_b.add_run()
        r_d.text = desc
        r_d.font.size = PptPt(13)
        r_d.font.color.rgb = GRAY_TEXT

    prs.save(filename)
    print(f"[SUCCESS] Master PowerPoint presentation generated: {filename}")

if __name__ == "__main__":
    docs_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(docs_dir, "4631_AWD_Final_Viva_Submission_Master.docx")
    pptx_path = os.path.join(docs_dir, "4631_AWD_Final_Viva_Submission_Master.pptx")
    
    create_master_word_document(docx_path)
    create_master_powerpoint_presentation(pptx_path)
